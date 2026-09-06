from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


class ExcelHandler:

    @staticmethod
    def unmerge_and_fill(input_file: str | Path,
                         output_file: str | Path) -> None:
        """
        Unmerge all merged cells and fill each cell in the merged range
        with the value of the top-left cell.
        """

        workbook = load_workbook(input_file)

        for sheet in workbook.worksheets:

            # Copy the merged ranges because we'll modify them
            merged_ranges = list(sheet.merged_cells.ranges)

            for merged_range in merged_ranges:

                min_row = merged_range.min_row
                max_row = merged_range.max_row
                min_col = merged_range.min_col
                max_col = merged_range.max_col

                # Get the value from the top-left cell
                value = sheet.cell(row=min_row, column=min_col).value

                # Unmerge the cells
                sheet.unmerge_cells(str(merged_range))

                # Fill all cells with the same value
                for row in range(min_row, max_row + 1):
                    for col in range(min_col, max_col + 1):
                        sheet.cell(row=row, column=col).value = value

        workbook.save(output_file)

    @staticmethod
    def read_excel(file_path: str | Path) -> pd.DataFrame:
        """
        Read an Excel file and return it as a DataFrame.
        """
        return pd.read_excel(file_path)

    @staticmethod
    def write_excel(dataframe: pd.DataFrame,
                    output_file: str | Path) -> None:
        """
        Save a DataFrame to an Excel file.
        """
        dataframe.to_excel(output_file, index=False)